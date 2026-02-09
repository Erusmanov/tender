import os
import zipfile
import re
from pathlib import Path

try:
    from PyPDF2 import PdfReader
    PYPDF2_AVAILABLE = True
except ImportError:
    PYPDF2_AVAILABLE = False

try:
    import pdfplumber
    PDFPLUMBER_AVAILABLE = True
except ImportError:
    PDFPLUMBER_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from bs4 import BeautifulSoup
    BS4_AVAILABLE = True
except ImportError:
    BS4_AVAILABLE = False

class DocumentReader:
    def __init__(self):
        self.supported_formats = ['txt', 'rtf', 'zip']
        if PYPDF2_AVAILABLE or PDFPLUMBER_AVAILABLE:
            self.supported_formats.append('pdf')
        if DOCX_AVAILABLE:
            self.supported_formats.extend(['docx', 'doc'])
        if BS4_AVAILABLE:
            self.supported_formats.append('html')

    def read_file(self, filepath: str) -> str:
        path = Path(filepath)
        if not path.exists():
            return f"[Error] File not found: {filepath}"
        
        ext = path.suffix.lower().lstrip('.')
        if not ext:
            if self._is_html_content(filepath):
                return self._read_html(filepath)
            return self._read_txt(filepath)

        method_name = f'_read_{ext}'
        if hasattr(self, method_name):
            return getattr(self, method_name)(filepath)
        
        if ext in ['htm', 'html']:
            return self._read_html(filepath)
        if ext == 'doc':
            return self._read_docx(filepath)
            
        return f"[Unsupported format] {ext}: {filepath}"

    def read_directory(self, directory: str, recursive: bool = True) -> str:
        dir_path = Path(directory)
        if not dir_path.exists():
            return ""

        files = list(dir_path.rglob('*')) if recursive else list(dir_path.glob('*'))
        files = sorted([f for f in files if f.is_file() and not f.name.startswith('.')])
        
        results = []
        results.append(f"=== Directory Content: {dir_path.name} ===")
        
        for filepath in files:
            ext = filepath.suffix.lower().lstrip('.')
            if ext not in self.supported_formats and ext != 'zip' and ext:
                continue

            content = self.read_file(str(filepath))
            if len(content) > 10000:
                content = content[:10000] + "\n[... truncated ...]"
            
            results.append(f"\n{'='*40}\nFile: {filepath.relative_to(dir_path)}\n{'='*40}")
            results.append(content)
            
        return "\n".join(results)

    def _read_txt(self, filepath: str) -> str:
        encodings = ['utf-8', 'cp1251', 'utf-16', 'iso-8859-5']
        for encoding in encodings:
            try:
                with open(filepath, 'r', encoding=encoding) as f:
                    return f.read()
            except (UnicodeDecodeError, UnicodeError):
                continue
        return ""

    def _read_pdf(self, filepath: str) -> str:
        text = []
        if PDFPLUMBER_AVAILABLE:
            try:
                with pdfplumber.open(filepath) as pdf:
                    text = [p.extract_text() or "" for p in pdf.pages]
                return "\n".join(text).strip()
            except Exception:
                pass

        if PYPDF2_AVAILABLE:
            try:
                reader = PdfReader(filepath)
                text = [p.extract_text() or "" for p in reader.pages]
                return "\n".join(text).strip()
            except Exception as e:
                return str(e)
        return ""

    def _read_docx(self, filepath: str) -> str:
        if not DOCX_AVAILABLE:
            return ""
        try:
            doc = Document(filepath)
            paragraphs = [p.text for p in doc.paragraphs]
            tables = []
            for table in doc.tables:
                for row in table.rows:
                    tables.append(" | ".join([c.text.strip() for c in row.cells]))
            return "\n".join(paragraphs + tables).strip()
        except Exception as e:
            return str(e)

    def _read_rtf(self, filepath: str) -> str:
        try:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
            text = re.sub(r'\\a-z+\\d*', '', content)
            text = re.sub(r'[{}]', '', text)
            return re.sub(r'\\.[^ ]+', '', text).strip()
        except Exception as e:
            return str(e)

    def _read_html(self, filepath: str) -> str:
        if not BS4_AVAILABLE:
            return ""
        try:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                soup = BeautifulSoup(f.read(), 'html.parser')
            
            for tag in soup(['script', 'style', 'meta', 'link']):
                tag.decompose()
            
            text_parts = []
            for tag in soup.find_all(['p', 'h1', 'h2', 'h3', 'div', 'li']):
                text = tag.get_text(separator=' ', strip=True)
                if len(text) > 3:
                    text_parts.append(text)
            
            for table in soup.find_all('table'):
                rows = []
                for tr in table.find_all('tr'):
                    cells = [td.get_text(strip=True) for td in tr.find_all(['td', 'th'])]
                    if cells:
                        rows.append(' | '.join(cells))
                if rows:
                    text_parts.append('\n'.join(rows))

            return '\n'.join(text_parts)
        except Exception as e:
            return str(e)

    def _read_zip(self, filepath: str) -> str:
        try:
            result = [f"=== ZIP Archive: {Path(filepath).name} ==="]
            with zipfile.ZipFile(filepath, 'r') as z:
                for name in z.namelist():
                    if name.startswith('__MACOSX') or name.endswith('/'):
                        continue
                    if any(name.endswith(e) for e in ['.txt', '.xml', '.json', '.md']):
                        try:
                            content = z.read(name).decode('utf-8', errors='ignore')
                            result.append(f"---\n{content[:5000]}")
                        except Exception:
                            pass
            return "\n\n".join(result)
        except Exception as e:
            return str(e)

    def _is_html_content(self, filepath: str) -> bool:
        try:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                head = f.read(500).lower()
                return any(x in head for x in ['<html', '<!doctype html', '<body'])
        except Exception:
            return False